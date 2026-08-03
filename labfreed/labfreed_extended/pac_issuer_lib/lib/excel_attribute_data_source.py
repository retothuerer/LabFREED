from __future__ import annotations

import logging
from datetime import datetime
from typing import Optional, Tuple, List, Dict

from cachetools import TTLCache, cachedmethod

from labfreed.pac_attributes.api_data_models.response import Spec_AttributeGroup
from labfreed.pac_attributes.facade.py_attributes import Attribute, Attributes
from labfreed.pac_attributes.server.server import AttributeGroupDataSource

try:
    from openpyxl import load_workbook
except ImportError:
    raise ImportError("Please install labfreed with the [extended] extra: pip install labfreed[extended]")

# ---------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------
def _get_row_by_first_cell(sheet_rows: List[tuple], match_value: str, base_url: str) -> Optional[Dict[str, object]]:
    if not sheet_rows:
        return None
    headers = sheet_rows[0]
    for row in sheet_rows[1:]:
        if not row:
            continue
        first = str(row[0]).strip() if row[0] is not None else ""
        if first == match_value:
            return {
                base_url + str(headers[i]).strip(): row[i]
                for i in range(1, len(headers))
                if headers[i] is not None
            }
    return None


# ---------------------------------------------------------------------
# Base class
# ---------------------------------------------------------------------
class _BaseExcelAttributeDataSource(AttributeGroupDataSource):
    """
    Common mapping logic from Excel rows to Spec_AttributeGroup.
    Subclasses implement `_read_rows_and_last_changed()`.
    """

    def __init__(self, *, base_url: str = "", cache_duration_seconds: int = 0, header_mappings=None, **kwargs):
        self._base_url = base_url
        self._header_mappings = header_mappings or dict()
        # a cache of this instance's own, so a different cache_duration_seconds on
        # another instance can never affect this one (see design-choices.md, "Each
        # Excel data source instance gets its own TTLCache")
        try:
            ttl = int(cache_duration_seconds)
        except Exception:
            ttl = 0
        self._cache = TTLCache(maxsize=128, ttl=ttl)
        super().__init__(**kwargs)

    def is_static(self) -> bool:
        return False

    def _read_rows_and_last_changed(self) -> Tuple[List[tuple], Optional[datetime]]:
        raise NotImplementedError

    @property
    def provides_attributes(self) -> List[str]:
        rows, _ = self._read_rows_and_last_changed()
        if not rows:
            return []
        return [self._base_url + r for r in rows[0][1:]]

    def attributes(self, subject_id:str) -> Optional[Spec_AttributeGroup]:
        rows, last_changed = self._read_rows_and_last_changed()

        d = None
        for key in self._lookup_keys(subject_id):
            d = _get_row_by_first_cell(rows, key, self._base_url)
            if d:
                break
        if not d:
            return None
        attributes = [Attribute(key= self._header_mappings.get(k, k), value=v) for k, v in d.items() if v is not None]
        return Spec_AttributeGroup(
            group_key=self.attribute_group_key,
            attributes=Attributes(attributes).to_payload_attributes()
        )


# ---------------------------------------------------------------------
# Local file implementation
# ---------------------------------------------------------------------
class LocalExcelAttributeDataSource(_BaseExcelAttributeDataSource):
    def __init__(self, file_path: str, **kwargs):
        self._file_path = file_path
        super().__init__(**kwargs)

    @cachedmethod(lambda self: self._cache)
    def _read_rows_and_last_changed(self) -> Tuple[List[tuple], Optional[datetime]]:
        logging.info(f"Attempting to load workbook: {self._file_path!r}")

        try:
            wb = load_workbook(
                filename=self._file_path,
                read_only=True,
                data_only=True
            )
            ws = wb.active
            logging.info(f"Workbook opened successfully. Active sheet: {ws.title!r}")

            rows = list(ws.iter_rows(values_only=True))
            logging.info(f"Read {len(rows)} rows from {self._file_path!r}")

            last_changed = wb.properties.modified
            logging.info(f"Workbook 'modified' property: {last_changed}")

            wb.close()
            return rows, last_changed

        except FileNotFoundError:
            logging.error(f"Workbook not found at: {self._file_path!r}", exc_info=True)
            raise
        except PermissionError:
            logging.error(f"Permission denied when accessing: {self._file_path!r}", exc_info=True)
            raise
        except Exception as e:
            logging.exception(f"Unexpected error reading workbook {self._file_path!r}: {e}")
            raise


__all__ = [
    "LocalExcelAttributeDataSource",
]
