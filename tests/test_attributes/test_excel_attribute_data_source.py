from unittest.mock import patch

import pytest
from openpyxl import Workbook

from labfreed.labfreed_extended.pac_issuer_lib.lib.excel_attribute_data_source import (
    LocalExcelAttributeDataSource,
    _BaseExcelAttributeDataSource,
    _get_row_by_first_cell,
)

TRAILING_SLASH_PAC_ID = "HTTPS://PAC.METTORIUS.COM/-MD/BAL500/12345/"


def _write_workbook(path, key):
    wb = Workbook()
    ws = wb.active
    ws.append(["id", "MfgDate"])
    ws.append([key, "2024-01-01"])
    wb.save(path)


def test_local_excel_data_source_finds_a_trailing_slash_id(tmp_path):
    # same regression class already fixed in Dict_DataSource.attributes(): a trailing-
    # slash id is now a *valid* PAC-ID, so PAC_CAT.from_url(...).to_url() canonicalizes
    # it and drops the slash - if the sheet is keyed by the original, uncanonicalized
    # string (as real-world data often is), a canonical-only lookup misses even though
    # the row is right there. Must fall back to the raw input, same as Dict_DataSource.
    file_path = tmp_path / "attributes.xlsx"
    _write_workbook(file_path, TRAILING_SLASH_PAC_ID)

    source = LocalExcelAttributeDataSource(
        file_path=str(file_path),
        attribute_group_key="instrument",
        cache_duration_seconds=0,
    )
    group = source.attributes(TRAILING_SLASH_PAC_ID)

    assert group is not None
    assert "MfgDate" in group.attributes


def test_local_excel_data_source_does_not_crash_on_a_non_pac_id(tmp_path):
    # a subject_id only has to be an IRI (see design-choices.md, "PAC-ID-Attributes'
    # subject_id accepts any IRI..."); today's bare `except:` happens to swallow this,
    # but the fix should catch LabFREED_ValidationError specifically, not any exception.
    non_pac_iri = "https://example.com/thing?x=1"
    file_path = tmp_path / "attributes.xlsx"
    _write_workbook(file_path, non_pac_iri)

    source = LocalExcelAttributeDataSource(
        file_path=str(file_path),
        attribute_group_key="instrument",
        cache_duration_seconds=0,
    )
    group = source.attributes(non_pac_iri)

    assert group is not None
    assert "MfgDate" in group.attributes


def test_two_instances_with_different_ttls_dont_stomp_each_others_cache(tmp_path):
    # regression test: _BaseExcelAttributeDataSource used to mutate one shared,
    # module-level TTLCache's .ttl in __init__, so constructing a second instance
    # with a different cache_duration_seconds silently changed the *first*
    # instance's caching too, since they shared the exact same cache object.
    file_a = tmp_path / "a.xlsx"
    _write_workbook(file_a, "KEY_A")

    source_a = LocalExcelAttributeDataSource(
        file_path=str(file_a),
        attribute_group_key="a",
        cache_duration_seconds=1000,
    )
    source_a.attributes("KEY_A")  # populate source_a's cache

    # constructing a second instance with a short TTL must not shorten source_a's
    # own, already-populated cache.
    LocalExcelAttributeDataSource(
        file_path=str(tmp_path / "b.xlsx"),
        attribute_group_key="b",
        cache_duration_seconds=0,
    )

    with patch(
        "labfreed.labfreed_extended.pac_issuer_lib.lib.excel_attribute_data_source.load_workbook"
    ) as mock_load:
        source_a.attributes("KEY_A")
    mock_load.assert_not_called()  # still within source_a's own long TTL, no re-read


class _FakeExcelSource(_BaseExcelAttributeDataSource):
    """Stands in for a real Excel-backed source in tests that only need to drive
    _BaseExcelAttributeDataSource's own mapping logic, without touching openpyxl."""

    def __init__(self, rows, **kwargs):
        self._rows = rows
        super().__init__(**kwargs)

    def _read_rows_and_last_changed(self):
        return self._rows, None


def test_get_row_by_first_cell_returns_none_for_empty_sheet():
    assert _get_row_by_first_cell([], "KEY", "") is None


def test_get_row_by_first_cell_skips_falsy_rows():
    rows = [("id", "MfgDate"), (), ("KEY", "val")]
    assert _get_row_by_first_cell(rows, "KEY", "") == {"MfgDate": "val"}


def test_base_class_defaults_ttl_to_zero_when_cache_duration_is_not_numeric():
    source = _BaseExcelAttributeDataSource(attribute_group_key="g", cache_duration_seconds=None)
    assert source._cache.ttl == 0


def test_base_class_is_static_is_false():
    source = _BaseExcelAttributeDataSource(attribute_group_key="g")
    assert source.is_static() is False




def test_provides_attributes_is_empty_when_there_are_no_rows():
    source = _FakeExcelSource(rows=[], attribute_group_key="g")
    assert source.provides_attributes == []


def test_provides_attributes_lists_header_row_after_the_key_column():
    source = _FakeExcelSource(rows=[("id", "A")], attribute_group_key="g")
    assert source.provides_attributes == ["A"]


def test_attributes_returns_none_when_no_row_matches_the_subject_id():
    source = _FakeExcelSource(rows=[("id", "A"), ("nomatch", "v")], attribute_group_key="g")
    assert source.attributes("KEY_NOT_PRESENT") is None


def test_local_excel_data_source_reraises_file_not_found_error(tmp_path):
    source = LocalExcelAttributeDataSource(
        file_path=str(tmp_path / "missing.xlsx"),
        attribute_group_key="g",
    )
    with patch(
        "labfreed.labfreed_extended.pac_issuer_lib.lib.excel_attribute_data_source.load_workbook",
        side_effect=FileNotFoundError("no such file"),
    ):
        with pytest.raises(FileNotFoundError):
            source._read_rows_and_last_changed()


def test_local_excel_data_source_reraises_permission_error(tmp_path):
    source = LocalExcelAttributeDataSource(
        file_path=str(tmp_path / "locked.xlsx"),
        attribute_group_key="g",
    )
    with patch(
        "labfreed.labfreed_extended.pac_issuer_lib.lib.excel_attribute_data_source.load_workbook",
        side_effect=PermissionError("denied"),
    ):
        with pytest.raises(PermissionError):
            source._read_rows_and_last_changed()


def test_local_excel_data_source_reraises_unexpected_errors(tmp_path):
    source = LocalExcelAttributeDataSource(
        file_path=str(tmp_path / "corrupt.xlsx"),
        attribute_group_key="g",
    )
    with patch(
        "labfreed.labfreed_extended.pac_issuer_lib.lib.excel_attribute_data_source.load_workbook",
        side_effect=ValueError("boom"),
    ):
        with pytest.raises(ValueError):
            source._read_rows_and_last_changed()
